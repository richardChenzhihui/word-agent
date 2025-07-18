"""
word_editor/test_docx_utils_full.py

紧凑测试集：用最少操作验证 docx_utils.py 的全部主要功能。
"""
import sys
import os
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))
from docx_utils import *
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.section import WD_ORIENT

def main():
    doc = create_new_docx()
    # 文档属性
    set_doc_property(doc, 'author', 'AI Tester')
    set_doc_property(doc, 'title', '功能测试文档')

    # 段落/标题/分页/列表
    add_heading(doc, '一级标题', level=1)
    p1 = add_paragraph(doc, '正文段落1')
    p2 = add_paragraph(doc, '正文段落2')
    insert_paragraph_before(p2, '插入段落在前')
    insert_paragraph_after(p1, '插入段落在后')
    add_page_break(doc)
    add_list(doc, ['无序1', '无序2'], ordered=False)
    add_list(doc, ['有序1', '有序2'], ordered=True)

    # 字体与样式
    run = p1.add_run(' 粗体红色')
    set_run_style(run, font_name='Arial', font_size=16, bold=True, color='#FF0000')
    set_paragraph_style(p1, alignment=WD_PARAGRAPH_ALIGNMENT.CENTER, line_spacing=2, indent=20)

    # 表格
    table = add_table(doc, 2, 2, style='Table Grid')
    set_cell_text(table.cell(0,0), '表格A')
    set_cell_text(table.cell(0,1), '表格B')
    set_cell_text(table.cell(1,0), '表格C')
    set_cell_text(table.cell(1,1), '表格D')
    merge_cells(table, 0, 0, 1, 0)
    set_table_style(table, 'Light Shading')

    # 图片
    # 需准备一张图片 word_editor/test_img.png
    img_path = os.path.join(os.path.dirname(__file__), 'test_img.png')
    if os.path.exists(img_path):
        add_picture(doc, img_path, width=1.0)

    # 超链接
    para_link = add_paragraph(doc, '访问: ')
    add_hyperlink(para_link, 'https://www.python.org', 'Python官网')

    # 页眉页脚/节
    header = get_header(doc)
    header.add_paragraph('这是页眉')
    footer = get_footer(doc)
    footer.add_paragraph('这是页脚')
    set_section_properties(doc.sections[0], orientation=WD_ORIENT.LANDSCAPE, margin=(100000,100000,100000,100000))

    # 高级文本替换
    add_paragraph(doc, '替换测试：hello world! hello world!')
    replace_text_in_docx(doc, 'hello', 'HELLO', nth=None)
    for table in doc.tables:
        replace_text_in_tables(doc, '表格', 'Table', nth=None)

    # 保存
    out_path = 'word_editor/test_docx_utils_full_output.docx'
    save_docx(doc, out_path)
    print(f'已生成完整功能测试文档: {out_path}\n请人工检查所有功能效果。')

if __name__ == '__main__':
    main() 