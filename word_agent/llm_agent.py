# llm_word_agent.py  — 轻量优化 & 稳定版（含表格 / 单元格自动解析）
import json, inspect, textwrap, re, time
from typing import Dict, List, Any, Optional, Tuple

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table, _Cell  # ← 新增 _Cell

# import docx_utils as du  # Word 操作函数集合
import docx_utils as du

# --------------------------  LLMConfig  -------------------------- #
class LLMConfig:
    def __init__(
            self,
            api_key: str,
            model: str = "gpt-3.5-turbo",
            base_url: Optional[str] = None,
            api_type: str = "openai",
            api_version: Optional[str] = None,
            temperature: float = 0.2,
            max_tokens: int = 2000,
            organization: Optional[str] = None,
    ):
        self.api_key, self.model = api_key, model
        self.base_url, self.api_type, self.api_version = base_url, api_type, api_version
        self.temperature, self.max_tokens, self.organization = temperature, max_tokens, organization


# =============================  Agent  ============================= #
class LLMWordAgent:
    def __init__(self, llm_config: Optional[LLMConfig] = None):
        self.llm_config = llm_config
        self.tools_meta = self._scan_tools()  # {name: (func, sig, intro)}
        self.utils_doc = self._build_tools_doc()
        self.runtime_ctx = {"last_table": None}  # ← 运行期上下文
        self.token_usage = {"total_tokens": 0, "prompt_tokens": 0, "completion_tokens": 0}
        self.progress_callback = None  # GUI进度回调
        self._setup_llm_client()

    def set_progress_callback(self, callback):
        """设置进度回调函数，用于GUI显示进度"""
        self.progress_callback = callback

    def _log_progress(self, message: str, is_error: bool = False):
        """记录进度信息"""
        if self.progress_callback:
            self.progress_callback(message, is_error)
        else:
            print(f"{'[ERROR]' if is_error else '[INFO]'} {message}")

    # ---------- 扫描 docx_utils ---------- #
    def _scan_tools(self) -> Dict[str, tuple]:
        tools = {}
        for name, fn in vars(du).items():
            if name.startswith("_") or not inspect.isfunction(fn):
                continue
            sig = inspect.signature(fn)
            doc_str = (fn.__doc__ or "").strip()
            intro = doc_str.splitlines()[0] if doc_str else ""
            tools[name] = (fn, sig, intro)
        return tools

    def _build_tools_doc(self) -> str:
        lines = ["【Word 编辑工具 - 自动生成】"]
        for name, (_, sig, intro) in self.tools_meta.items():
            lines.append(f"{name}{sig}  –  {intro}")
        return "\n".join(lines)

    # ---------- LLM 客户端 ---------- #
    def _setup_llm_client(self):
        if not self.llm_config:
            self.client = None
            return

        import openai

        params = {"api_key": self.llm_config.api_key}
        if self.llm_config.organization:
            params["organization"] = self.llm_config.organization
        if self.llm_config.base_url:
            if self.llm_config.api_type == "azure":
                params.update(
                    azure_endpoint=self.llm_config.base_url,
                    api_version=self.llm_config.api_version or "2024-02-15-preview",
                )
            else:
                params["base_url"] = self.llm_config.base_url
        self.client = openai.OpenAI(**params)

    # ---------- Prompt ---------- #
    def _build_prompt(self, structure: Dict[str, Any], user_inst: str) -> str:
        # 检测是否为批量替换任务
        batch_replace_keywords = ['改写', '替换', '全部', '批量', '所有的', '改为', '换成']
        is_batch_replace = any(keyword in user_inst for keyword in batch_replace_keywords)
        
        batch_replace_guidance = ""
        if is_batch_replace:
            batch_replace_guidance = """
            
            【批量替换特别指导】
            检测到批量替换任务，请特别注意：
            1. 使用 batch_replace 函数进行批量文本替换，格式：{"action": "batch_replace", "replace_map": {"旧文本1": "新文本1", "旧文本2": "新文本2"}}
            2. 或使用多个 replace_text_in_docx 和 replace_text_in_tables 调用，确保覆盖文档和表格中的所有文本
            3. 对于占位符替换（如 [NAME OF OFC]），必须完整匹配，包括方括号
            4. 替换操作要放在其他操作之前，确保所有文本都被正确替换
            5. 如果有多个需要替换的文本，为每个文本创建单独的替换操作
            """
        
        return textwrap.dedent(
            f"""
            你是专业 Word 编辑助手。请根据用户指令，输出 docx_utils 调用指令。
            【可用工具】
            {self.utils_doc}

            【当前文档结构】
            {json.dumps(structure, ensure_ascii=False, indent=2)}

            【用户指令】
            {user_inst}
            {batch_replace_guidance}

            【输出要求】
            1. 仅输出 JSON 数组，每项包含 "action" 及对应参数，字段名与函数签名一致。
            2. 若函数首参为 doc（Document 对象）系统会自动注入，JSON 中不要提供 doc。
            3. Paragraph / Table 参数用 index（int）或 {{ "index": idx }} 表示。
            4. 如需单元格，可仅提供 {{ "row": r, "col": c }}，系统自动定位到最近创建的表格。
            5. 如果忘记提供 table 参数，系统默认使用最近一次 add_table 创建的表格。
            6. 对于批量替换任务，优先使用 batch_replace 或多个 replace_text_in_docx/replace_text_in_tables 操作。
            7. 不要输出除 JSON 之外的任何文本。
            """
        )

    # ---------- 调用 LLM ---------- #
    def call_llm(self, prompt: str) -> Tuple[List[Dict[str, Any]], float]:
        """调用LLM并返回命令列表和执行时间"""
        if not self.client:  # 离线演示
            # 智能检测批量替换任务
            if self._is_batch_replace_task(prompt):
                return self._generate_batch_replace_commands(prompt), 0.5
            else:
                return [{"action": "add_heading", "text": "示例文档", "level": 1}], 0.5

        start_time = time.time()
        
        try:
            rsp = self.client.chat.completions.create(
                model=self.llm_config.model,
                messages=[
                    {"role": "system", "content": "你是 Word 编辑助手，仅输出 JSON"},
                    {"role": "user", "content": prompt},
                ],
                temperature=self.llm_config.temperature,
                max_tokens=self.llm_config.max_tokens,
                response_format={"type": "json_object"},
            )
            
            # 记录token使用情况
            if hasattr(rsp, 'usage') and rsp.usage:
                self.token_usage["total_tokens"] += rsp.usage.total_tokens
                self.token_usage["prompt_tokens"] += rsp.usage.prompt_tokens
                self.token_usage["completion_tokens"] += rsp.usage.completion_tokens
            
            execution_time = time.time() - start_time
            content = rsp.choices[0].message.content
            
            self._log_progress(f"LLM响应时间: {execution_time:.2f}秒")
            self._log_progress(f"Token使用: {rsp.usage.total_tokens if hasattr(rsp, 'usage') and rsp.usage else 'N/A'}")
            
            # 解析 JSON
            try:
                data = json.loads(content)
                commands = data if isinstance(data, list) else data.get("commands", [])
                return commands, execution_time
            except json.JSONDecodeError:
                m = re.search(r"\[.*\]", content, re.S)
                if m:
                    return json.loads(m.group(0)), execution_time
                raise ValueError("LLM 输出无法解析为 JSON")
                
        except Exception as e:
            execution_time = time.time() - start_time
            self._log_progress(f"LLM调用失败: {str(e)}", is_error=True)
            raise e

    # ---------- 参数校验 & 转换 ---------- #
    def _validate_and_prepare(self, doc: Document, cmd: Dict[str, Any]):
        action = cmd.get("action")
        if action not in self.tools_meta:
            raise ValueError(f"未知 action: {action}")

        fn, sig, _ = self.tools_meta[action]
        kwargs = {k: v for k, v in cmd.items() if k != "action"}

        # -------- 始终注入 / 覆盖 doc -------- #
        for p in sig.parameters.values():
            if p.annotation is Document or p.name == "doc":
                kwargs[p.name] = doc

        # -------- 特殊处理 set_cell_text 的 row/col 参数 -------- #
        if action == "set_cell_text" and "row" in kwargs and "col" in kwargs:
            # 从 row/col 构造 cell 参数
            row = kwargs.pop("row")
            col = kwargs.pop("col")
            table_index = kwargs.pop("table_index", None)
            
            if table_index is not None:
                table = doc.tables[table_index]
            elif self.runtime_ctx.get("last_table"):
                table = self.runtime_ctx["last_table"]
            else:
                raise ValueError("无法确定表格：缺少 table_index 且没有最近创建的表格")
            
            kwargs["cell"] = table.cell(row, col)

        # -------- 将 index / dict 转对象 -------- #
        def to_para(v):
            if isinstance(v, int):
                return doc.paragraphs[v]
            if isinstance(v, dict):
                return doc.paragraphs[v.get("index", 0)]
            return v

        def to_table(v):
            if isinstance(v, int):
                return doc.tables[v]
            if isinstance(v, dict):
                return doc.tables[v.get("index", 0)]
            # v 为空 → 使用最近表格（若存在）
            if v is None and self.runtime_ctx.get("last_table"):
                return self.runtime_ctx["last_table"]
            return v

        def to_cell(v):
            if isinstance(v, _Cell):
                return v
            if isinstance(v, dict) and {"row", "col"} <= v.keys():
                tbl = (
                    doc.tables[v.get("table_index")]
                    if "table_index" in v
                    else to_table(None)
                )
                return tbl.cell(v["row"], v["col"])
            return v

        # 遍历函数参数，自动转换
        for p in sig.parameters.values():
            name_lower = p.name.lower()
            ann = p.annotation

            # 若参数缺失但需要 Table，对其使用最近表格
            if p.name not in kwargs and (
                    ann is Table or name_lower == "table"
            ) and self.runtime_ctx.get("last_table"):
                kwargs[p.name] = self.runtime_ctx["last_table"]

            if p.name not in kwargs:
                continue

            val = kwargs[p.name]

            if ann is Paragraph or (
                    ann is inspect._empty and "para" in name_lower
            ):
                kwargs[p.name] = to_para(val)

            elif ann is Table or (
                    ann is inspect._empty and "table" in name_lower
            ):
                kwargs[p.name] = to_table(val)

            elif ann is _Cell or (
                    ann is inspect._empty and "cell" in name_lower
            ):
                kwargs[p.name] = to_cell(val)

        # -------- 必填 / 多余检查 -------- #
        for p in sig.parameters.values():
            if p.default is p.empty and p.name not in kwargs:
                raise ValueError(f"缺少必需参数 {p.name} for {action}")
        extra = set(kwargs) - set(sig.parameters)
        if extra:
            raise ValueError(f"多余参数 {extra} for {action}")

        return fn, kwargs

    # ---------- 执行指令 ---------- #
    def execute_commands(self, doc: Document, cmds: List[Dict[str, Any]]):
        res = {"success": [], "errors": []}
        for idx, cmd in enumerate(cmds, 1):
            try:
                fn, kwargs = self._validate_and_prepare(doc, cmd)
                out = fn(**kwargs)  # 真正调用

                # 若执行 add_table，记录返回表格
                if cmd["action"] == "add_table" and isinstance(out, Table):
                    self.runtime_ctx["last_table"] = out

                success_msg = f"#{idx} {cmd['action']} 执行成功"
                res["success"].append(success_msg)
                self._log_progress(success_msg)
            except Exception as e:
                error_msg = f"#{idx} {cmd['action']}: {str(e)}"
                res["errors"].append(error_msg)
                self._log_progress(error_msg, is_error=True)
        return res

    # ---------- 重试逻辑 ---------- #
    def _should_retry(self, exec_res: Dict, execution_time: float) -> bool:
        """判断是否需要重试"""
        # 执行时间小于2秒
        if execution_time < 2.0:
            return True
        
        # 有错误且没有成功操作
        if exec_res["errors"] and not exec_res["success"]:
            return True
            
        # 没有任何成功操作
        if not exec_res["success"]:
            return True
            
        return False

    def _create_document_copy(self, original_doc: Document, docx_path: str) -> Document:
        """创建文档副本用于测试"""
        import io
        import tempfile
        
        # 如果是空文档，直接创建新的
        if not docx_path or docx_path.strip() == "":
            return du.create_new_docx()
        
        # 否则，通过保存和重新加载来创建副本
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp_file:
            original_doc.save(tmp_file.name)
            return du.open_docx(tmp_file.name)
    
    def _execute_with_retry(self, doc: Document, structure: Dict[str, Any], user_inst: str, max_attempts: int = 15, docx_path: str = "") -> Dict:
        """带重试逻辑的执行"""
        attempts = []
        perfect_attempt_found = False  # 标记是否找到了完美执行（成功>0且错误=0）
        
        self._log_progress(f"开始执行，最多尝试 {max_attempts} 次")
        
        for attempt in range(max_attempts):
            try:
                self._log_progress(f"第 {attempt + 1} 次尝试...")
                
                # 为每次尝试创建文档副本
                test_doc = self._create_document_copy(doc, docx_path)
                
                # 重置运行时上下文
                self.runtime_ctx = {"last_table": None}
                
                # 构建prompt并调用LLM
                prompt = self._build_prompt(structure, user_inst)
                commands, execution_time = self.call_llm(prompt)
                
                # 在测试文档上执行命令
                exec_res = self.execute_commands(test_doc, commands)
                
                # 记录这次尝试的结果
                attempt_result = {
                    "attempt": attempt + 1,
                    "commands": commands,
                    "execution_results": exec_res,
                    "execution_time": execution_time,
                    "success_count": len(exec_res["success"]),
                    "error_count": len(exec_res["errors"]),
                    "test_doc": test_doc  # 保存测试文档以备后用
                }
                attempts.append(attempt_result)
                
                self._log_progress(f"第 {attempt + 1} 次尝试完成: 成功 {len(exec_res['success'])} 个, 错误 {len(exec_res['errors'])} 个")
                
                # 检查是否找到了完美执行
                if len(exec_res["success"]) > 0 and len(exec_res["errors"]) == 0:
                    perfect_attempt_found = True
                    self._log_progress(f"发现完美执行！成功 {len(exec_res['success'])} 个操作，0 个错误")
                    break
                
                # 检查是否需要重试（只有在没有找到完美执行时才继续重试）
                if not self._should_retry(exec_res, execution_time):
                    break
                        
            except Exception as e:
                error_msg = f"第 {attempt + 1} 次尝试失败: {str(e)}"
                self._log_progress(error_msg, is_error=True)
                
                attempt_result = {
                    "attempt": attempt + 1,
                    "commands": [],
                    "execution_results": {"success": [], "errors": [error_msg]},
                    "execution_time": 0,
                    "success_count": 0,
                    "error_count": 1,
                    "test_doc": None
                }
                attempts.append(attempt_result)
        
        # 如果找到了完美执行，再尝试3次以寻找更好的结果
        if perfect_attempt_found and len(attempts) < max_attempts:
            remaining_attempts = min(3, max_attempts - len(attempts))
            self._log_progress(f"已找到完美执行，再尝试 {remaining_attempts} 次以寻找更优结果...")
            
            for extra_attempt in range(remaining_attempts):
                try:
                    attempt_num = len(attempts) + 1
                    self._log_progress(f"优化第 {extra_attempt + 1} 次尝试 (总第 {attempt_num} 次)...")
                    
                    # 为每次尝试创建文档副本
                    test_doc = self._create_document_copy(doc, docx_path)
                    
                    # 重置运行时上下文
                    self.runtime_ctx = {"last_table": None}
                    
                    # 构建prompt并调用LLM
                    prompt = self._build_prompt(structure, user_inst)
                    commands, execution_time = self.call_llm(prompt)
                    
                    # 在测试文档上执行命令
                    exec_res = self.execute_commands(test_doc, commands)
                    
                    # 记录这次尝试的结果
                    attempt_result = {
                        "attempt": attempt_num,
                        "commands": commands,
                        "execution_results": exec_res,
                        "execution_time": execution_time,
                        "success_count": len(exec_res["success"]),
                        "error_count": len(exec_res["errors"]),
                        "test_doc": test_doc
                    }
                    attempts.append(attempt_result)
                    
                    self._log_progress(f"优化第 {extra_attempt + 1} 次尝试完成: 成功 {len(exec_res['success'])} 个, 错误 {len(exec_res['errors'])} 个")
                    
                except Exception as e:
                    error_msg = f"优化第 {extra_attempt + 1} 次尝试失败: {str(e)}"
                    self._log_progress(error_msg, is_error=True)
                    
                    attempt_result = {
                        "attempt": len(attempts) + 1,
                        "commands": [],
                        "execution_results": {"success": [], "errors": [error_msg]},
                        "execution_time": 0,
                        "success_count": 0,
                        "error_count": 1,
                        "test_doc": None
                    }
                    attempts.append(attempt_result)
        
        # 选择成功操作最多的一次作为最终结果
        if attempts:
            best_attempt = max(attempts, key=lambda x: x["success_count"])
            self._log_progress(f"选择第 {best_attempt['attempt']} 次尝试作为最终结果 (成功 {best_attempt['success_count']} 个操作)")
            
            # 将最佳结果应用到原始文档
            if best_attempt["test_doc"] and best_attempt["success_count"] > 0:
                self._log_progress("将最佳结果应用到最终文档...")
                # 使用最佳尝试的测试文档替换原始文档的内容
                self._copy_document_content(best_attempt["test_doc"], doc)
                self._log_progress("最佳结果已应用到最终文档")
            
            # 输出所有请求消耗的token统计
            total_token_summary = f"总Token消耗: {self.token_usage['total_tokens']}, 其中输入Token: {self.token_usage['prompt_tokens']}, 输出Token: {self.token_usage['completion_tokens']}"
            self._log_progress(total_token_summary)
            
            return {
                "status": "success" if best_attempt["success_count"] > 0 else "partial_success",
                "best_attempt": best_attempt,
                "all_attempts": attempts,
                "total_attempts": len(attempts),
                "token_usage": self.token_usage,
                "perfect_attempt_found": perfect_attempt_found
            }
        else:
            return {
                "status": "error",
                "error": "所有尝试都失败了",
                "all_attempts": attempts,
                "total_attempts": len(attempts),
                "token_usage": self.token_usage,
                "perfect_attempt_found": False
            }
    
    def _copy_document_content(self, source_doc: Document, target_doc: Document):
        """将源文档的内容复制到目标文档"""
        # 清空目标文档
        for element in target_doc.element.body:
            target_doc.element.body.remove(element)
        
        # 复制源文档的内容
        for element in source_doc.element.body:
            target_doc.element.body.append(element)

    # ---------- 顶层接口 ---------- #
    def edit_document(self, docx_path: str, user_inst: str, out_path: str):
        """编辑文档的主要接口"""
        try:
            # 处理空路径情况 - 创建新文档
            if not docx_path or docx_path.strip() == "":
                self._log_progress("创建新的空白文档")
                doc = du.create_new_docx()
            else:
                self._log_progress(f"打开文档: {docx_path}")
                doc = du.open_docx(docx_path)
            
            # 获取文档结构
            structure = du.export_docx_structure(doc)
            self._log_progress(f"文档结构解析完成，包含 {len(structure.get('blocks', []))} 个块")
            
            # 重置token使用统计
            self.token_usage = {"total_tokens": 0, "prompt_tokens": 0, "completion_tokens": 0}
            
            # 执行带重试的编辑
            result = self._execute_with_retry(doc, structure, user_inst, docx_path=docx_path)
            
            # 保存文档
            self._log_progress(f"保存文档到: {out_path}")
            du.save_docx(doc, out_path)
            
            result["output"] = out_path
            return result
            
        except Exception as e:
            error_msg = f"编辑文档时发生错误: {str(e)}"
            self._log_progress(error_msg, is_error=True)
            return {
                "status": "error",
                "error": error_msg,
                "token_usage": self.token_usage
            }

    def get_token_usage(self) -> Dict[str, int]:
        """获取token使用统计"""
        return self.token_usage.copy()
    
    def _is_batch_replace_task(self, prompt: str) -> bool:
        """检测是否为批量替换任务"""
        batch_replace_keywords = ['改写', '替换', '全部', '批量', '所有的', '改为', '换成', '[NAME OF OFC]', '[Name of Sub-Fund]']
        return any(keyword in prompt for keyword in batch_replace_keywords)
    
    def _generate_batch_replace_commands(self, prompt: str) -> List[Dict[str, Any]]:
        """生成批量替换命令"""
        commands = []
        
        # 从prompt中提取替换映射
        replace_map = self._extract_replace_map(prompt)
        
        if replace_map:
            # 使用批量替换命令
            commands.append({
                "action": "batch_replace",
                "replace_map": replace_map
            })
        else:
            # 如果无法提取，使用默认的占位符替换
            default_replace_map = {
                "[NAME OF OFC]": "king talia",
                "[Name of Sub-Fund]": "queen richard"
            }
            commands.append({
                "action": "batch_replace",
                "replace_map": default_replace_map
            })
        
        return commands
    
    def _extract_replace_map(self, prompt: str) -> Dict[str, str]:
        """从prompt中提取替换映射"""
        replace_map = {}
        
        # 查找形如 "[NAME OF OFC] 改为 king talia" 的模式
        import re
        
        # 模式1: [占位符] 改为 新值
        pattern1 = r'\[([^\]]+)\]\s*(?:改为|换成|替换为|改写为|是)\s*([^\n\r]+)'
        matches1 = re.findall(pattern1, prompt, re.IGNORECASE)
        
        for placeholder, replacement in matches1:
            full_placeholder = f"[{placeholder}]"
            replace_map[full_placeholder] = replacement.strip()
        
        # 模式2: 占位符 是 新值
        pattern2 = r'\[([^\]]+)\]\s*是\s*([^\n\r]+)'
        matches2 = re.findall(pattern2, prompt, re.IGNORECASE)
        
        for placeholder, replacement in matches2:
            full_placeholder = f"[{placeholder}]"
            replace_map[full_placeholder] = replacement.strip()
        
        return replace_map


# ------------------------------ 示例 ------------------------------ #
if __name__ == "__main__":
    """
    直接运行此文件可执行一次“多功能综合测试”，需：
    • demo.docx 作为输入
    • finance.png（或自行修改图片路径）位于同级目录
    • 配置自己的 OpenAI 兼容 key
    """
    import os

    # os.environ["OPENAI_API_KEY"] = "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"

    cfg = LLMConfig(
        api_key=os.getenv("OPENAI_API_KEY"),
        base_url="https://openrouter.ai/api/v1",
        model="moonshotai/kimi-k2",
        temperature=0.7,
        # max_tokens=2048,
    )
    agent = LLMWordAgent(cfg if cfg.api_key else None)

    # user_inst = """
    # 请按以下顺序对文档进行多步骤编辑：
    #
    # 1) 在文档开头插入一级标题“示例文档”。
    # 2) 紧跟其后插入自我介绍段落：“大家好，我是大型语言模型驱动的 Word 智能助手，下面演示多功能编辑。”
    # 3) 在自我介绍段落后插入分页符。
    # 4) 在新页添加二级标题“香港金融发展展望”。
    # 5) 在标题后插入三段文字（你生成），讨论香港金融业未来 5 年的主要趋势。
    # 6) 然后插入一个无序列表，内容为：绿色金融、离岸人民币、金融科技。
    # 7) 接着插入一个 3×3 表格：
    #    - 第一行（标题行）：指标 | 2023 | 2024 （并将这一整行 3 个单元格合并）
    #    - 第二行：IPO数量 | 90 | 120
    #    - 第三行：融资总额 | 1200亿 | 1500亿
    #    - 表格样式设为 “Table Grid”。
    # 8) 把正文和表格中出现的“香港”全部替换为“H.K.”（保留格式）。
    # 9) 在表格后再插入分页符并插入图片 finance.png，宽 4 英寸。
    # 10) 在文档末尾插入一个有序列表，列出你认为香港金融可持续增长的 3 个关键词。
    # 11) 设置文档作者属性为 “LLM-Agent”。
    # 12) 把所有的文字替换成英文翻译，你（大模型）提供翻译
    # """

    #
    #
    #
    # user_inst = """
    # Please read the existing contract text in the current document and replace each occurrence of the old contract elements listed below with the corresponding new values. Do not modify any other clauses, numbering, or formatting—only perform precise text substitutions. Return **only** a JSON array of `docx_utils` commands.
    #
    # New contract elements to apply:
    # - Party A: Beijing ZhiXin Technology Co., Ltd.
    # - Party B: Shanghai Internet Network Co., Ltd.
    # - Signing Date: July 1, 2025
    # - Performance Location: Wangjing Street, Chaoyang District, Beijing
    # - Total Contract Amount: RMB 800,000.00
    #
    # And change all text into time new roman
    # """
    #
    # res = agent.edit_document(
    #     docx_path="demo.docx",
    #     user_inst=user_inst,
    #     out_path="demo_out.docx",
    # )
    # print(json.dumps(res, ensure_ascii=False, indent=2))
    #
    # user_inst = """
    # Please read the existing legal text in the current document and replace each occurrence of the old contract elements listed below with the corresponding new values. Do not modify any other clauses, numbering, or formatting—only perform precise text substitutions. Return **only** a JSON array of `docx_utils` commands.
    #
    # New contract elements to apply:
    # initial limited partner 的名字和信息如下：姓名：秦易，身份证地址：四川省广安市广安区龙台镇竹埝村3组12号。
    # input text should be the english translation, you should read the document and find all places need modification
    # """
    #
    # res = agent.edit_document(
    #     docx_path="./documents/1. Sierra Education LPF - Initial LPA - Execution Version.docx",
    #     user_inst=user_inst,
    #     out_path="./documents/1. Sierra Education LPF - Initial LPA - Execution Version_out.docx",
    # )
    # print(json.dumps(res, ensure_ascii=False, indent=2))

    # 演示新功能
    user_inst = """
    创建一个现代化的技术文档，包含以下内容：
    1. 添加标题"Word Agent 功能演示"
    2. 创建一个包含重试机制、空文档支持、Token统计等新功能的介绍段落
    3. 插入一个表格展示新功能对比
    4. 将所有字体改为宋体
    5. 在文档末尾添加版本信息
    """

    # 演示空文档创建
    print("演示1: 创建新文档")
    res = agent.edit_document(
        docx_path="",  # 空路径，将创建新文档
        user_inst=user_inst,
        out_path="documents/demo_new_features.docx",
    )
    print(json.dumps(res, ensure_ascii=False, indent=2))
    
    # 演示Token统计
    print(f"\nToken使用统计: {agent.get_token_usage()}")
    
    print("\n" + "="*50)
    print("GUI已准备就绪！运行以下命令启动图形界面：")
    print("python run_gui.py")
    print("或者：")
    print("python word_agent_gui.py")
    # #

    # 已经发现的bug
    # 1. 原本的docx 不能是空的否则无法读取
    # 2. add_list 有问题
    # 3. 律所文件出现发现执行修改后，word 没有变动的地方
