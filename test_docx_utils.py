"""
word_editor/test_docx_utils.py

测试 docx_utils.py 的保留格式精确替换功能，包括：
- 普通段落替换
- 多次/全部替换
- 保留格式（如粗体、斜体）
- 表格内替换
- 边界情况
"""
import sys
import os
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))
from docx import Document
from docx_utils import (
    create_new_docx, save_docx, add_paragraph, add_table,
    replace_text_in_docx, replace_text_in_tables
)

def test_paragraph_replace():
    doc = create_new_docx()
    add_paragraph(doc, "Hello world! This is a test.")
    add_paragraph(doc, "Another line with test word.")
    add_paragraph(doc, "No match here.")
    path = "word_editor/test_para.docx"
    save_docx(doc, path)
    # 替换第一个 test
    doc2 = Document(path)
    count = replace_text_in_docx(doc2, "test", "exam", nth=1)
    save_docx(doc2, "word_editor/test_para_replaced1.docx")
    print(f"[段落] 替换第1个 test: 替换次数={count}, 输出=word_editor/test_para_replaced1.docx")
    # 替换全部 test
    doc3 = Document(path)
    count = replace_text_in_docx(doc3, "test", "exam", nth=None)
    save_docx(doc3, "word_editor/test_para_replaced_all.docx")
    print(f"[段落] 替换全部 test: 替换次数={count}, 输出=word_editor/test_para_replaced_all.docx")

def test_format_preserve():
    doc = create_new_docx()
    p = add_paragraph(doc, "This is a ")
    run1 = p.add_run("bold")
    run1.bold = True
    p.add_run(" and ")
    run2 = p.add_run("italic.")
    run2.italic = True
    path = "word_editor/test_format.docx"
    save_docx(doc, path)
    doc2 = Document(path)
    count = replace_text_in_docx(doc2, "bold", "BOLD", nth=None)
    save_docx(doc2, "word_editor/test_format_replaced.docx")
    print(f"[格式] 替换 bold: 替换次数={count}, 输出=word_editor/test_format_replaced.docx")
    # 检查 BOLD 是否仍为粗体，italic 是否仍为斜体

def test_table_replace():
    doc = create_new_docx()
    table = add_table(doc, 2, 2)
    table.cell(0,0).text = "cell test"
    table.cell(0,1).text = "no match"
    table.cell(1,0).text = "test in table"
    table.cell(1,1).text = "test test"
    path = "word_editor/test_table.docx"
    save_docx(doc, path)
    doc2 = Document(path)
    count = replace_text_in_tables(doc2, "test", "exam", nth=None)
    save_docx(doc2, "word_editor/test_table_replaced.docx")
    print(f"[表格] 替换全部 test: 替换次数={count}, 输出=word_editor/test_table_replaced.docx")

def test_edge_cases():
    doc = create_new_docx()
    add_paragraph(doc, "")  # 空段落
    add_paragraph(doc, "test")  # 只有目标词
    add_paragraph(doc, "testtest")  # 连续目标词
    add_paragraph(doc, "t e s t")  # 不应匹配
    path = "word_editor/test_edge.docx"
    save_docx(doc, path)
    doc2 = Document(path)
    count = replace_text_in_docx(doc2, "test", "exam", nth=None)
    save_docx(doc2, "word_editor/test_edge_replaced.docx")
    print(f"[边界] 替换全部 test: 替换次数={count}, 输出=word_editor/test_edge_replaced.docx")

def main():
    test_paragraph_replace()
    test_format_preserve()
    test_table_replace()
    test_edge_cases()
    print("\n请手动检查 word_editor 目录下的输出文件，确认替换效果和格式保留情况。")

if __name__ == "__main__":
    main() 